# -*- coding: utf-8 -*-
"""
export_docx.py — 将 paper/论文.md 转为 paper/论文.docx（python-docx 实现）
支持：标题层级、表格、图片、粗体、行内代码、列表、代码块。
运行：python code/export_docx.py
"""
import re, sys
from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

WORK = Path(__file__).resolve().parent.parent
SRC = WORK / "paper" / "论文.md"
DST = WORK / "paper" / "论文.docx"


def set_cn_font(run, size=None, bold=None, mono=False):
    run.font.name = "Consolas" if mono else "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体" if not mono else "Consolas")
    if size: run.font.size = Pt(size)
    if bold is not None: run.font.bold = bold


def split_md_row(line):
    """按未转义的 | 切分表格行，并把 \\| 还原为 |（保证数学式内竖线不串列）。"""
    s = line.strip()
    if s.startswith("|"):
        s = s[1:]
    if s.endswith("|"):
        s = s[:-1]
    parts = re.split(r"(?<!\\)\|", s)
    return [c.replace("\\|", "|").strip() for c in parts]


def add_rich(par, text):
    """解析 **bold** 与 `code`。"""
    pos = 0
    for m in re.finditer(r"\*\*(.+?)\*\*|`([^`]+)`", text):
        if m.start() > pos:
            set_cn_font(par.add_run(text[pos:m.start()]))
        if m.group(1) is not None:
            r = par.add_run(m.group(1)); set_cn_font(r, bold=True)
        else:
            r = par.add_run(m.group(2)); set_cn_font(r, mono=True)
        pos = m.end()
    if pos < len(text):
        set_cn_font(par.add_run(text[pos:]))


def main():
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    style.font.size = Pt(11)

    lines = SRC.read_text(encoding="utf-8").splitlines()
    i, in_code = 0, False
    code_buf = []
    while i < len(lines):
        ln = lines[i]
        if ln.strip().startswith("```"):
            if in_code:
                p = doc.add_paragraph()
                r = p.add_run("\n".join(code_buf)); set_cn_font(r, size=9, mono=True)
                code_buf = []
            in_code = not in_code
            i += 1; continue
        if in_code:
            code_buf.append(ln); i += 1; continue
        if ln.startswith("|") and i + 1 < len(lines) and re.match(r"^\|[\s:\-|]+\|$", lines[i+1].strip()):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(split_md_row(lines[i])); i += 1
            rows.pop(1)  # 分隔行
            ncol = max(len(r) for r in rows)
            t = doc.add_table(rows=len(rows), cols=ncol)
            t.style = "Table Grid"
            for ri, row in enumerate(rows):
                for ci in range(ncol):
                    cell = t.cell(ri, ci)
                    cell.text = ""
                    p = cell.paragraphs[0]
                    add_rich(p, row[ci] if ci < len(row) else "")
                    for r in p.runs:
                        r.font.size = Pt(9.5)
                        if ri == 0: r.font.bold = True
            doc.add_paragraph()
            continue
        m = re.match(r"^(#{1,4})\s+(.*)$", ln)
        if m:
            level = len(m.group(1))
            h = doc.add_heading("", level=min(level, 4) if level > 1 else 1)
            r = h.add_run(m.group(2))
            r.font.color.rgb = RGBColor(0, 0, 0)
            r.font.name = "Times New Roman"
            r._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
            i += 1; continue
        m = re.match(r"^!\[[^\]]*\]\(([^)]+)\)", ln.strip())
        if m:
            img = (WORK / "paper" / m.group(1)).resolve()
            if img.exists():
                doc.add_picture(str(img), width=Cm(15.5))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            i += 1; continue
        s = ln.strip()
        if s.startswith("- "):
            p = doc.add_paragraph(style="List Bullet"); add_rich(p, s[2:])
        elif re.match(r"^\d+\.\s", s):
            # 手工编号普通段落（不用 Word List Number 样式，避免跨章节连续编号）
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.74)
            p.paragraph_format.first_line_indent = Cm(-0.74)
            add_rich(p, s)
        elif s == "---":
            pass
        elif s:
            p = doc.add_paragraph(); add_rich(p, s)
        i += 1

    doc.save(DST)
    print(f"已导出 {DST}")


if __name__ == "__main__":
    main()
